#!/usr/bin/env python3
"""Update the startup appendix while preserving the existing v10 guide pages.

Requires reportlab, pypdf and python-pptx. Run from any directory. The original
22 pages/slides remain untouched; a prior appendix produced here is replaced.
The PDF and editable PPTX share the same text, dimensions and coordinates.
"""
from io import BytesIO
from pathlib import Path
import argparse

from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
NAME = "M5Stack_Inventory_技術者向け構築・検証ガイド_v10"
TITLE = "StickS3の起動確認と切り分け"
WIDTH, HEIGHT = 960, 540
EMU = 12700
FONT_NAME = "Hiragino Sans"
INK, MUTED, BLUE, WHITE = "132333", "5C7182", "356FF0", "FFFFFF"

# Positions use PDF points, measured from the upper-left corner.
TEXTS = [
    (42, 24, 876, 16, BLUE, "付録  StickS3 / Unit QRCode"),
    (42, 60, 876, 28, INK, TITLE),
    (42, 111, 876, 13, MUTED, "Unit起動時のブートモードから、電源を切り替えず通常の読取りへ戻すための機能です。"),
    (42, 156, 440, 18, BLUE, "READYになるまで"),
    (42, 192, 440, 13, INK, "① 初期化中もIRを消灯（GPIO46 LOW）\n   最初の給電後800msは通信せず待つ\n   正常0x21の設定を確認してREADYへ"),
    (42, 255, 440, 13, INK, "② ブート0x54だけが応答した場合\n   復帰命令0x77を一度だけ送る\n   電源は切り替えない"),
    (42, 318, 440, 13, INK, "③ 最大約2秒・10回で0x21だけを確認\n   版1〜254・手動1・TRIG 0/1まで\n   読戻しが正常ならREADYへ"),
    (42, 387, 440, 11, MUTED, "起動・通信・復帰の段階を115200bpsで記録。\nコード本文・Wi-Fi情報・キーは記録しません。"),
    (522, 156, 396, 18, BLUE, "確認中と停止時の表示"),
    (522, 192, 396, 13, INK, "STARTING / READER RESUME：復帰確認中\nERROR / RECOVERY HOLD：復帰未確認\nREADER STORAGE：復帰の記録保存に失敗\nREADER POWER / BUS / I2C：給電・通信\nREADER CONFIG：版・設定・TRIG読戻し"),
    (522, 300, 396, 12, MUTED, "復帰命令の前に確認待ちを端末へ保存。\n失敗や本体再起動後も0x54を再確認せず、\n命令を再送しません。正常0x21の設定まで\n確認できた場合だけ記録を解除します。\nAボタンで再試行せず、ログを確認します。\n未送信イベントは保持。全消去しないでください。"),
    (42, 441, 876, 12, INK, "ブート復帰は実機未検証です。Unit起動時の不調を対象とし、スキャン中の異常とは区別します。\n電源原因の確定や再発防止の実証とは扱わず、READY・実読取り・Scans保存を別々に確認します。"),
    (42, 516, 280, 10, MUTED, "M5Stack 在庫管理システム"),
    (365, 516, 430, 10, MUTED, "© 2026 K Visualization Studio"),
    (878, 516, 40, 10, MUTED, "23"),
]


def build_pdf(pdf_path, font_path):
    pdfmetrics.registerFont(TTFont("JP", str(font_path)))
    buffer = BytesIO()
    c = canvas.Canvas(buffer, pagesize=(WIDTH, HEIGHT), pageCompression=1)
    c.setFillColor(HexColor("#F4F7FA"))
    c.rect(0, 0, WIDTH, HEIGHT, fill=1, stroke=0)
    c.setFillColor(HexColor("#00B8E6"))
    c.rect(0, 0, 12, HEIGHT, fill=1, stroke=0)
    c.setStrokeColor(HexColor("#D9E3EA"))
    c.line(42, 40, 918, 40)
    for x, top, width, size, color, value in TEXTS:
        c.setFillColor(HexColor("#" + color))
        c.setFont("JP", size)
        for i, line in enumerate(value.splitlines()):
            if pdfmetrics.stringWidth(line, "JP", size) > width:
                raise ValueError(f"Text exceeds its box: {line}")
            c.drawString(x, HEIGHT - top - size - i * size * 1.45, line)
    c.save()
    appendix = PdfReader(buffer)
    original = PdfReader(pdf_path)
    if len(original.pages) not in (22, 23):
        raise ValueError("Expected the original 22 pages and optional startup appendix")
    if len(original.pages) == 23 and TITLE not in (original.pages[-1].extract_text() or ""):
        raise ValueError("Refusing to overwrite an unrelated 23rd page")
    output = PdfWriter()
    for page in original.pages[:22]:
        output.add_page(page)
    output.add_page(appendix.pages[0])
    output.add_metadata({"/Title": "M5Stack在庫管理システム 技術者向け構築・検証ガイド", "/Author": "K Visualization Studio"})
    with pdf_path.open("wb") as stream:
        output.write(stream)


def build_pptx(pptx_path):
    deck = Presentation(pptx_path)
    if len(deck.slides) not in (22, 23):
        raise ValueError("Expected 22 slides and optional startup appendix")
    if len(deck.slides) == 23:
        if TITLE not in "\n".join(s.text for s in deck.slides[-1].shapes if s.has_text_frame):
            raise ValueError("Refusing to overwrite an unrelated 23rd slide")
        slide_id = deck.slides._sldIdLst[-1]
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    if (deck.slide_width, deck.slide_height) != (WIDTH * EMU, HEIGHT * EMU):
        raise ValueError("Unexpected guide dimensions")
    slide = deck.slides.add_slide(deck.slides[-1].slide_layout)
    for placeholder in list(slide.placeholders):
        element = placeholder._element
        element.getparent().remove(element)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("F4F7FA")
    from pptx.enum.shapes import MSO_SHAPE
    for x, y, w, h, color in [(0, 0, 12, HEIGHT, "00B8E6"), (42, 500, 876, 1, "D9E3EA")]:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x * EMU, y * EMU, w * EMU, h * EMU)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
    for x, y, width, size, color, value in TEXTS:
        box = slide.shapes.add_textbox(x * EMU, y * EMU, width * EMU, int((value.count("\n") + 1) * size * 1.45 * EMU + 4 * EMU))
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        for i, line in enumerate(value.splitlines()):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = FONT_NAME
            p.font.size = Pt(size)
            p.font.color.rgb = RGBColor.from_string(color)
            p.font.underline = False
            p.line_spacing = 1.45
            p.space_before = p.space_after = Pt(0)
    deck.save(pptx_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font", type=Path, default=Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"))
    args = parser.parse_args()
    if not args.font.is_file():
        parser.error("Provide a TrueType font containing Japanese characters with --font")
    pdf_path = ROOT / "docs/M5Stack_Inventory_Guide.pdf"
    pptx_path = ROOT / "docs/M5Stack_Inventory_Guide.pptx"
    # Build both outputs before replacing either published artifact.
    import tempfile
    import shutil
    with tempfile.TemporaryDirectory() as temporary:
        pdf_candidate = Path(temporary) / pdf_path.name
        pptx_candidate = Path(temporary) / pptx_path.name
        shutil.copy2(pdf_path, pdf_candidate)
        shutil.copy2(pptx_path, pptx_candidate)
        build_pdf(pdf_candidate, args.font)
        build_pptx(pptx_candidate)
        shutil.copy2(pdf_candidate, pdf_path)
        shutil.copy2(pptx_candidate, pptx_path)
    print("Updated startup appendix: 23 PDF pages and 23 editable PPTX slides")


if __name__ == "__main__":
    main()
